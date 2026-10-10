"""Independent layout contracts and actual compatibility qualification, never paper templates."""
from copy import deepcopy
from pathlib import Path
from io import BytesIO
import hashlib
import json
import pytest
from lxml import etree as E
from fontTools.ttLib import TTFont
from word_module_studio.package import Package, W, M, NS, snapshot
from word_module_studio.context import LayoutContext
from word_module_studio.templates import build, check_fonts
from word_module_studio.runtime import WordError, run, digest, write_json
from word_module_studio.render import render
from test_templates import spec
from test_native_modules import apply, item


def test_effective_style_theme_and_explicit_false(document, tmp_path):
    output = apply(document, tmp_path, 'M19', item(document, 'document'), action='extension',
        props={'operation': 'theme', 'properties': {'theme.name': 'Layout test', 'theme.font.minor.latin': 'Liberation Sans', 'theme.font.minor.cs': 'Noto Sans Arabic', 'theme.color.accent1': '227788'}})
    package = Package(output); styles = package.xml('word/styles.xml')
    base = E.SubElement(styles, W + 'style'); base.set(W + 'type', 'paragraph'); base.set(W + 'styleId', 'LayoutBase')
    E.SubElement(base, W + 'name').set(W + 'val', 'LayoutBase')
    rp = E.SubElement(base, W + 'rPr'); E.SubElement(rp, W + 'rFonts').set(W + 'asciiTheme', 'minorHAnsi')
    pp = E.SubElement(base, W + 'pPr'); E.SubElement(pp, W + 'keepNext')
    leaf = E.SubElement(styles, W + 'style'); leaf.set(W + 'type', 'paragraph'); leaf.set(W + 'styleId', 'LayoutLeaf')
    E.SubElement(leaf, W + 'name').set(W + 'val', 'LayoutLeaf'); E.SubElement(leaf, W + 'basedOn').set(W + 'val', 'LayoutBase')
    target = next(o for o in package.objects() if o['kind'] == 'paragraph' and o['text'] == 'first')
    pp = target['element'].find(W + 'pPr')
    if pp is None: pp = E.Element(W + 'pPr'); target['element'].insert(0, pp)
    for n in list(pp): pp.remove(n)
    E.SubElement(pp, W + 'pStyle').set(W + 'val', 'LayoutLeaf'); E.SubElement(pp, W + 'keepNext').set(W + 'val', '0')
    context = LayoutContext.resolve(package, target)
    assert context.fonts['ascii'] == 'Liberation Sans'
    assert not context.pagination['keepNext']
    assert context.effective_format['sources']['font.asciiTheme'] == 'LayoutBase'
    package.save()
    proof = render(output, tmp_path / 'theme-proof')
    assert any('LiberationSans' in name.replace('-', '') for name in proof['actual_fonts']), proof['actual_fonts']
    E.SubElement(base, W + 'basedOn').set(W + 'val', 'LayoutLeaf')
    with pytest.raises(WordError, match='STYLE_INHERITANCE_CYCLE'): LayoutContext.resolve(package, target)


def test_header_and_note_ownership(tmp_path):
    policy = spec(3); policy['regions'][0]['header'] = 'First section header'
    policy['regions'][1]['header'] = 'Second section header'
    manifest = build(policy, tmp_path / 'template'); package = Package(manifest['master'])
    header = next(o for o in package.objects() if o['kind'] == 'paragraph' and o['text'] == 'First section header')
    context = LayoutContext.resolve(package, header)
    assert context.owning_sections == (1,) and context.container == 'page'
    assert context.available_width_twips == context.width_twips
    sections = list(package.xml('word/document.xml').iter(W + 'sectPr'))
    second = sections[1].find(W + 'headerReference'); sections[1].remove(second)
    assert LayoutContext.resolve(package, header).owning_sections == (1, 2)


def test_native_textbox_content_width(document, tmp_path):
    output = apply(document, tmp_path, 'M13', item(document, 'body'), action='add', kind='textbox',
        props={'text': 'Textbox width', 'width': '4cm', 'height': '2cm'})
    package = Package(output)
    target = next(o for o in package.objects() if o['kind'] == 'paragraph' and o['text'] == 'Textbox width' and any(E.QName(a).localname == 'txbxContent' for a in o['element'].iterancestors()))
    context = LayoutContext.resolve(package, target)
    assert context.container == 'textbox' and 1700 < context.available_width_twips < 2400


def test_native_frame_and_original_preserved(document, tmp_path):
    before = digest(document)
    output = apply(document, tmp_path, 'M03', item(document, 'paragraph', 'first'), action='extension',
        props={'operation': 'frame', 'attributes': {'dropCap': 'drop', 'lines': 3, 'hSpace': 120, 'wrap': 'around'}})
    node = Package(output).xml('word/document.xml').find('.//' + W + 'framePr')
    assert node.get(W + 'dropCap') == 'drop' and node.get(W + 'lines') == '3'
    assert digest(document) == before


def test_font_is_really_embedded_and_reversible(document, tmp_path):
    font = check_fonts({'latin': 'Liberation Serif'})['latin']
    output = apply(document, tmp_path, 'M19', item(document, 'document'), action='extension', props={
        'operation': 'embed_fonts', 'fonts': [{**{k: font[k] for k in ('family', 'file', 'index')}, 'redistributable': True}]})
    package = Package(output); table = package.xml('word/fontTable.xml')
    ref = table.find('.//' + W + 'embedRegular'); assert ref is not None
    rel = next(n for n in package.xml('word/_rels/fontTable.xml.rels') if n.get('Id') == ref.get('{' + NS['r'] + '}id'))
    import posixpath
    target = rel.get('Target')
    data = bytearray(package.parts[posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join('word', target))])
    mask = bytes.fromhex(ref.get(W + 'fontKey').strip('{}').replace('-', ''))[::-1]
    for index in range(32): data[index] ^= mask[index % 16]
    decoded = TTFont(BytesIO(data)); assert decoded['name'].getDebugName(1) == 'Liberation Serif'; decoded.close()
    report = json.loads(output.with_suffix('.operations.json').read_text())
    assert hashlib.sha256(data).hexdigest() == report['operations'][0]['result']['sources'][0]['embedded_sha256']


def test_restricted_font_and_missing_declaration_are_rejected(document, tmp_path):
    font = check_fonts({'latin': 'Liberation Serif'})['latin']; restricted = tmp_path / 'restricted.ttf'
    with TTFont(font['file']) as value: value['OS/2'].fsType = 2; value.save(restricted)
    for path, declared, error in ((restricted, True, 'FONT_EMBEDDING_RESTRICTED'), (Path(font['file']), False, 'FONT_DISTRIBUTION_DECLARATION_REQUIRED')):
        with pytest.raises(WordError, match=error):
            apply(document, tmp_path, 'M19', item(document, 'document'), action='extension', props={
                'operation': 'embed_fonts', 'fonts': [{'family': 'Liberation Serif', 'file': str(path), 'redistributable': declared}]})


def test_composable_typography_theme_table_and_math_settings(tmp_path):
    from word_module_studio.compose import compose
    policy = spec(2); policy['regions'] = [policy['regions'][1]]
    policy.update(typography={'docGrid.type': 'lines', 'docGrid.linePitch': 300, 'kinsoku': True, 'autoHyphenation': True},
        theme={'theme.color.accent1': '227788', 'theme.font.minor.latin': 'Liberation Serif'},
        table_policy={'layout': 'autofit', 'preferred_width_pct': 90},
        equations={'long_formula': 'practical', 'operator_position': 'before', 'continuation_indent_twips': 300})
    policy['styles']['Normal']['native'] = {'widowControl': True, 'kinsoku': True, 'keepLines': False}
    policy['regions'][0]['native'] = {'vAlign': 'top', 'marginHeader': 500}
    manifest = build(policy, tmp_path / 'template')
    markdown = tmp_path / 'generic.md'; markdown.write_text('# Layout core\n\n中文排版测试 Paragraph wrapping.\n\n| Name | Value |\n|---|---|\n| retained | 123 |', encoding='utf-8')
    output = tmp_path / 'generic.docx'; compose(markdown, manifest, output)
    package = Package(output)
    assert package.xml('word/document.xml').find('.//' + W + 'docGrid').get(W + 'linePitch') == '300'
    assert package.xml('word/document.xml').find('.//' + W + 'tblLayout').get(W + 'type') == 'autofit'
    assert package.xml('word/settings.xml').find('.//' + M + 'wrapIndent').get(M + 'val') == '300'
    report = render(output, tmp_path / 'proof'); assert report['automatic_pass'], report['issues']


def test_qualify_renderer_native_math_breaks(document, tmp_path):
    # A capability probe records true compatibility rather than calling unsupported output a pass.
    from word_module_studio.math import convert
    formula = convert('a+b+c+d+e+f+g+h+i+j+k+l+m+n+p+q+r+s+t+u+v+w+x+y+z')
    runs = list(formula.findall(M + 'r'))
    operators = [r for r in runs if (r.findtext(M + 't') or '').startswith('+')]
    if not operators: operators = [r for r in runs if '+' in (r.findtext(M + 't') or '')]
    if operators:
        target = operators[len(operators) // 2]
        properties = target.find(M + 'rPr')
        if properties is None: properties = E.Element(M + 'rPr'); target.insert(0, properties)
        E.SubElement(properties, M + 'brk')
    package = Package(document); body = package.xml('word/document.xml').find(W + 'body')
    p = E.Element(W + 'p'); display = E.SubElement(p, M + 'oMathPara'); display.append(formula); body.insert(0, p)
    output = tmp_path / 'native-break-probe.docx'; package.save(output)
    report = render(output, tmp_path / 'proof')
    import pdfplumber
    with pdfplumber.open(tmp_path / 'proof/native-break-probe.pdf') as pdf:
        page = pdf.pages[0]
        ys = sorted({round(c['top'] / 3) * 3 for c in page.chars if c['text'] == '+'})
    write_json(tmp_path / 'native-break-qualification.json', {'renderer': report['renderer'], 'operator_runs': len(operators),
        'native_manual_break_visible': len(ys) >= 2, 'operator_line_positions': ys, 'source_sha256': digest(output)})
    assert operators, 'Probe must place a real native operator break'


def test_practical_math_preserves_author_alignment(document, tmp_path):
    from word_module_studio.migration.semantic_import import normalize_aligned
    formula = r'\begin{aligned}x+y&=1\\x-y&=0\end{aligned}'
    assert normalize_aligned(formula) == formula
    output = apply(document, tmp_path, 'M09', item(document, 'body'), action='extension',
        props={'formula': formula, 'number': 'A1', 'policy': 'practical'})
    root = Package(output).xml('word/document.xml')
    assert ''.join(root.xpath('//m:t/text()', namespaces=NS)).count('=') == 2
    assert root.find('.//' + M + 'm') is not None or root.find('.//' + M + 'eqArr') is not None
    assert '&' not in root.xpath('//m:t/text()', namespaces=NS)
    report = render(output, tmp_path / 'alignment-proof'); assert report['automatic_pass'], report['issues']


def test_practical_math_never_splits_nested_atoms(document, tmp_path):
    from word_module_studio.math import convert
    numerator = '+'.join('x_{' + str(i) + '}' for i in range(35))
    formula = r'\frac{' + numerator + '}{1+x^2}'
    expected = convert(formula).find(M + 'f')
    output = apply(document, tmp_path, 'M09', item(document, 'body'), action='extension',
        props={'formula': formula, 'number': 'F1', 'policy': 'practical'})
    actual = Package(output).xml('word/document.xml').find('.//' + M + 'f')
    from word_module_studio.migration.math_layout import tree
    for node in actual.iter():
        for key in list(node.attrib):
            if key.startswith('{' + NS['wms'] + '}'): del node.attrib[key]
    assert tree(actual) == tree(expected)
    assert actual.find('.//' + M + 'eqArr') is None
    report = json.loads(output.with_suffix('.operations.json').read_text())['operations'][0]['result']
    assert report['layout_status'] == 'needs_minor_adjustment' and report['manual_adjustments']


def test_field_cache_refresh_changes_only_results(document, tmp_path):
    from word_module_studio.field_cache import refresh, fields, without_cache
    package = Package(document); body = package.xml('word/document.xml').find(W + 'body')
    p = E.Element(W + 'p'); field = E.SubElement(p, W + 'fldSimple'); field.set(W + 'instr', ' NUMPAGES ')
    E.SubElement(E.SubElement(field, W + 'r'), W + 't').text = '999'; body.insert(0, p)
    package.save(); before = digest(document); baseline = without_cache(Package(document))
    output = tmp_path / 'cached.docx'; result = refresh(document, output, tmp_path / 'cache-proof')
    assert result['fields_updated'][0]['result_text'] == '1'
    assert without_cache(Package(output)) == baseline and digest(document) == before
    assert result['cache_convergence_passes'] == 2


def _field(body, instruction, cached='STALE'):
    p = E.SubElement(body, W + 'p')
    for kind in ('begin', 'separate', 'end'):
        if kind == 'separate':
            E.SubElement(E.SubElement(p, W + 'r'), W + 'instrText').text = instruction
        E.SubElement(E.SubElement(p, W + 'r'), W + 'fldChar').set(W + 'fldCharType', kind)
        if kind == 'separate': E.SubElement(E.SubElement(p, W + 'r'), W + 't').text = cached
    return p


def test_toc_and_caption_index_refresh_after_edit(tmp_path):
    from docx import Document
    from docx.shared import Pt
    from word_module_studio.field_cache import refresh, fields, result_nodes, without_cache
    source = tmp_path / 'indexes.docx'; doc = Document()
    doc.styles['Normal'].font.name = 'Liberation Serif'; doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading('Index proof', 1)
    doc.add_heading('Retained chapter', 1); doc.add_paragraph('RETAINED section text.')
    doc.add_page_break(); doc.add_heading('Deleted chapter', 1); doc.add_paragraph('REMOVE section text.')
    doc.save(source); package = Package(source); body = package.xml('word/document.xml').find(W + 'body')
    toc = _field(body, r' TOC \o "1-3" \h '); body.remove(toc); body.insert(1, toc)
    caption = _field(body, r' SEQ Figure ', '1')
    E.SubElement(E.SubElement(caption, W + 'r'), W + 't').text = ' Caption evidence'
    index = _field(body, r' TOC \c "Figure" \h '); body.remove(index); body.insert(2, index)
    _field(body, r' XE "Retained topic" ', '')
    _field(body, r' INDEX \h "A" ')
    # Keep section properties last.
    section = body.find(W + 'sectPr'); body.remove(section); body.append(section)
    package.save(); before = digest(source); baseline = without_cache(Package(source))
    first = tmp_path / 'first.docx'; report = refresh(source, first, tmp_path / 'first-proof')
    assert digest(source) == before and without_cache(Package(first)) == baseline
    records = report['fields_updated']; toc_text = next(r['result_text'] for r in records if '\\o' in r['instruction'])
    assert 'Retained chapter' in toc_text and 'Deleted chapter' in toc_text
    assert any('Caption evidence' in r['result_text'] for r in records if '\\c' in r['instruction'])
    assert any('Retained topic' in r['result_text'] for r in records if r['instruction'].strip().startswith('INDEX'))
    doc = Document(first)
    for p in list(doc.paragraphs):
        if p.text == 'Deleted chapter' or p.text == 'REMOVE section text.': p._element.getparent().remove(p._element)
    doc.add_page_break(); doc.add_heading('Added chapter', 1); doc.add_paragraph('ADDED content.')
    edited = tmp_path / 'edited.docx'; doc.save(edited)
    final = tmp_path / 'final.docx'; report = refresh(edited, final, tmp_path / 'final-proof')
    text = next(r['result_text'] for r in report['fields_updated'] if '\\o' in r['instruction'])
    assert 'Deleted chapter' not in text and 'Added chapter' in text and 'Retained chapter' in text
    assert report['automatic_pass'], report['proof']['issues']


@pytest.mark.parametrize('variant', ['single', 'double', 'triple', 'unequal', 'landscape'])
def test_generic_layout_gallery(tmp_path, variant):
    from word_module_studio.compose import compose
    columns = {'single': 1, 'double': 2, 'triple': 3, 'unequal': 2, 'landscape': 2}[variant]
    policy = spec(columns); policy['id'] = 'core-' + variant
    policy['regions'] = [policy['regions'][1]]; region = policy['regions'][0]
    region.update(header='Core layout / ' + variant, footer='Page {{PAGE}} / {{NUMPAGES}}')
    if variant == 'unequal': region['column_widths_twips'] = [3500, 5572]
    if variant == 'landscape': region['page'].update(width_twips=16838, height_twips=11906)
    policy['fonts']['rtl'] = 'Noto Sans Arabic'
    policy.update(typography={'kinsoku': True, 'autoHyphenation': True},
        table_policy={'layout': 'autofit', 'preferred_width_pct': 95}, equations={'long_formula': 'practical'})
    policy['styles']['Normal'].update(native={'widowControl': True, 'keepLines': False})
    manifest = build(policy, tmp_path / 'template')
    paragraphs = ['# Core layout ' + variant, '中文排版，标点禁则与分页测试。中文内容保留。',
        'مرحبا بالعالم', '1. First numbered observation\n2. Second numbered observation']
    for index in range(8): paragraphs.append('REGION' + str(index) + ' ' + 'Shared native modules keep text in its assigned container. ' * 15)
    paragraphs += [r'$$\begin{aligned}x+y&=1\\x-y&=0\end{aligned}\tag{A1}$$',
        '$$s=' + '+'.join('x_{' + str(i) + '}' for i in range(24)) + r'\tag{L1}$$',
        '| Item | Result |\n|---|---|\n| DATA_A | 123 |\n| DATA_B | 456 |\n\nTable: Retained data',
        'Footnote target.[^note]\n\n[^note]: Retained footnote evidence.']
    markdown = tmp_path / 'layout.md'; markdown.write_text('\n\n'.join(paragraphs), encoding='utf-8')
    output = tmp_path / ('core-' + variant + '.docx'); composition = compose(markdown, manifest, output)
    output = apply(output, tmp_path, 'M03', item(output, 'paragraph', 'مرحبا بالعالم'), props={'direction': 'rtl', 'align': 'right'}, name='direction')
    report = render(output, tmp_path / 'proof')
    assert report['automatic_pass'], report['issues']
    assert report['container_proof']['paragraphs']
    import pdfplumber
    with pdfplumber.open(tmp_path / 'proof/direction.pdf') as pdf:
        text = ''.join(p.extract_text() or '' for p in pdf.pages)
        assert all('REGION' + str(i) in text for i in range(8))
        assert '123' in text and '456' in text and '(A1)' in text and '(L1)' in text
    package = Package(output)
    assert package.xml('word/footnotes.xml').find('.//' + W + 't') is not None
    assert not any('split-fraction' in str(e['changes']) for e in composition['equations'])


@pytest.mark.parametrize('chart_type', ['column', 'bar', 'line', 'pie', 'area', 'scatter'])
def test_core_chart_gallery(document, tmp_path, chart_type):
    output = apply(document, tmp_path, 'M12', item(document, 'paragraph', 'second'), action='add', kind='chart',
        props={'chartType': chart_type, 'data': 'Series:1,3,2', 'categories': 'A,B,C', 'title': chart_type + ' editable chart', 'width': '11cm', 'height': '7cm'})
    report = render(output, tmp_path / 'proof'); assert report['automatic_pass'], report['issues']
    assert report['container_proof']['shapes']


def test_core_object_gallery(tmp_path):
    from word_module_studio.compose import compose
    policy = spec(1); policy['regions'] = [policy['regions'][1]]
    policy.update(font_embedding={'enabled': True, 'redistributable': True, 'roles': ['latin']})
    manifest = build(policy, tmp_path / 'template')
    source = tmp_path / 'objects.md'; source.write_text('# Core objects\n\nContainer anchor text.\n\nFloating picture anchor.\n\nShape anchor.\n\nTextbox anchor.\n\n| State | Detail |\n|---|---|\n| One | Merged native data |\n| Two | Retained second row |\n\nNote target.[^note]\n\n[^note]: Object gallery native note.', encoding='utf-8')
    output = tmp_path / 'objects.docx'; compose(source, manifest, output)
    output = apply(output, tmp_path, 'M10', item(output, 'table'), action='extension',
        props={'operation': 'merge', 'row_start': 1, 'row_end': 1, 'column_start': 0, 'column_end': 1}, name='merged-cell')
    output = apply(output, tmp_path, 'M11', item(output, 'paragraph', 'Floating picture anchor.'), action='add', kind='picture',
        props={'path': str(Path('fixtures/upstream/word/pictures-logo.png').resolve()), 'width': '3cm', 'height': '1.5cm', 'anchor': True, 'wrap': 'square',
            'hRelative': 'column', 'vRelative': 'paragraph', 'hPosition': '0cm', 'vPosition': '0cm', 'alt': 'Retained logo'}, name='image')
    output = apply(output, tmp_path, 'M13', item(output, 'body'), action='add', kind='shape',
        props={'geometry': 'ellipse', 'width': '3cm', 'height': '1cm', 'fill': '88AAEE', 'anchor.x': '0cm', 'anchor.y': '0cm'}, name='shape')
    output = apply(output, tmp_path, 'M13', item(output, 'body'), action='add', kind='textbox',
        props={'text': 'Editable box content', 'width': '5cm', 'height': '1.5cm', 'anchor.x': '4cm', 'anchor.y': '0cm'}, name='box')
    # Independent native frame interface uses a complete framed paragraph.
    output = apply(output, tmp_path, 'M03', item(output, 'paragraph', 'Container anchor text.'), action='extension',
        props={'operation': 'frame', 'attributes': {'w': 3000, 'h': 600, 'wrap': 'around', 'hAnchor': 'margin', 'vAnchor': 'text', 'hSpace': 120}}, name='frame')
    report = render(output, tmp_path / 'proof'); assert report['automatic_pass'], report['issues']
    assert report['container_proof']['frames'] or report['container_proof']['shapes']


def test_drop_cap_visual_gallery(document, tmp_path):
    output = apply(document, tmp_path, 'M03', item(document, 'paragraph', 'first'), props={'text': 'Drop cap paragraph demonstrates a native frame with retained editable text.'})
    output = apply(output, tmp_path, 'M03', item(output, 'paragraph', 'Drop cap paragraph demonstrates a native frame with retained editable text.'), action='extension',
        props={'operation': 'frame', 'attributes': {'dropCap': 'drop', 'lines': 3, 'hSpace': 120, 'wrap': 'around'}}, name='drop-cap')
    report = render(output, tmp_path / 'proof'); assert report['automatic_pass'], report['issues']
    import pdfplumber
    with pdfplumber.open(tmp_path / 'proof/drop-cap.pdf') as pdf:
        capitals = [c['size'] for p in pdf.pages for c in p.chars if c['text'] == 'D']
        body = [c['size'] for p in pdf.pages for c in p.chars if c['text'] == 's']
    assert capitals and body and max(capitals) > max(body) * 1.5, 'Native drop cap must appear at its requested line height'
