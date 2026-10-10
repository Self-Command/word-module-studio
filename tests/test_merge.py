from pathlib import Path
from docx import Document
from word_module_studio.package import Package, NS, W, snapshot
from word_module_studio.merge import merge, extract
from word_module_studio.backend import OfficeBackend


def make(path, name):
    doc = Document(); doc.add_paragraph(name); doc.add_paragraph('Shared section', style='Heading 1')
    doc.sections[0].header.paragraphs[0].text = name + ' header'
    doc.save(path)
    api = OfficeBackend()
    api.edit(path, '/body/p[1]', 'add', 'footnote', {'text': name + ' footnote'})
    api.edit(path, '/body/p[1]', 'add', 'endnote', {'text': name + ' endnote'})
    api.edit(path, '/body/p[1]', 'add', 'comment', {'text': name + ' comment', 'author': name})
    api.edit(path, '/body/p[2]', 'add', 'bookmark', {'name': 'SharedHeading'})


def test_merge_native_auxiliary_parts_and_source_headers(tmp_path):
    first, second, output = [tmp_path / n for n in ('first.docx', 'second.docx', 'joined.docx')]
    make(first, 'First'); make(second, 'Second')
    merge(first, [second], output, header_policy='source_regions')
    package = Package(output)
    for part in ('footnotes', 'endnotes', 'comments'):
        root = package.xml('word/' + part + '.xml')
        text = ''.join(root.xpath('//w:t/text()', namespaces=NS))
        assert 'First' in text and 'Second' in text
        ids = [n.get(W + 'id') for n in root]
        assert len(ids) == len(set(ids))
    root = package.xml('word/document.xml')
    names = [n.get(W + 'name') for n in root.iter(W + 'bookmarkStart')]
    assert len(names) == len(set(names))
    assert len(list(root.iter(W + 'sectPr'))) == 2
    headers = ''.join(package.parts[n].decode() for n in package.parts if n.startswith('word/header') and n.endswith('.xml'))
    assert 'First header' in headers and 'Second header' in headers
    from word_module_studio.render import render
    report = render(output, tmp_path / 'merge-proof')
    assert report['automatic_pass'], report['issues']


def test_extract_object_keeps_native_dependencies(document, tmp_path):
    api = OfficeBackend(); api.edit(document, '/body/p[1]', 'add', 'footnote', {'text': 'Selected note'})
    target = next(o for o in snapshot(document)['objects'] if o['kind'] == 'paragraph' and o['text'] == 'first')
    output = tmp_path / 'selected.docx'; extract(document, [target['id']], output)
    package = Package(output)
    assert 'second' not in ''.join(package.xml('word/document.xml').xpath('//w:t/text()', namespaces=NS))
    assert 'Selected note' in ''.join(package.xml('word/footnotes.xml').xpath('//w:t/text()', namespaces=NS))
