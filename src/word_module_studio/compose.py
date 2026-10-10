"""Markdown semantics enter registered regions; source geometry never does."""
from copy import deepcopy
import json
from pathlib import Path
import re
import tempfile
from urllib.parse import unquote, urlparse
from lxml import etree as E
from docx import Document
from .runtime import require_actions, WordError, run, write_json, digest
from .package import Package, W, M, NS
from .templates import check_identity, verify_locked
from .content_import import import_resources, import_lists
from .merge import NativeComposer

MARKDOWN = 'markdown+tex_math_single_backslash+tex_math_dollars+pipe_tables+table_captions+citations+raw_attribute-native_divs-native_spans'


def parse(path, resources, folder):
    path = Path(path).resolve()
    ast = json.loads(run(['pandoc', '-f', MARKDOWN, '-t', 'json', path], reject_warnings=True))
    from .migration.semantic_import import expand_semantic_tex
    ast = expand_semantic_tex(ast, folder)
    permitted = {Path(p).resolve() for p in resources}
    equations = []
    ignored_layout = []
    def visit(value):
        if isinstance(value, dict):
            if value.get('t') in ('RawInline', 'RawBlock'): raise WordError('CONTENT_RAW_OBJECT_FORBIDDEN')
            if value.get('t') == 'Image':
                url = value['c'][2][0]
                if urlparse(url).scheme: raise WordError('RESOURCE_MUST_BE_EXPLICIT_LOCAL_FILE', url)
                image = (path.parent / unquote(url)).resolve()
                if image not in permitted or not image.is_file(): raise WordError('IMAGE_RESOURCE_MISSING', url)
                value['c'][2][0] = str(image)
            # Pandoc attributes can express physical styles, so reject these before conversion.
            if value.get('t') in ('Image', 'Header', 'Div', 'Span', 'CodeBlock', 'Table'):
                attr = value['c'][1] if value['t'] == 'Header' else value['c'][0]
                if isinstance(attr, list) and len(attr) == 3:
                    if value['t'] == 'Image' and not attr[1] and all(key in ('width', 'height') for key, _ in attr[2]):
                        if attr[2]: ignored_layout.append({'kind': 'image', 'attributes': attr[2], 'reason': 'Registered region controls physical size'})
                        attr[2] = []
                    elif value['t'] == 'CodeBlock' and not attr[2] and len(attr[1]) <= 1:
                        pass  # Language identifier is semantic content, not a document style.
                    elif attr[1] or attr[2]: raise WordError('CONTENT_LAYOUT_ATTRIBUTE_FORBIDDEN', str(attr))
            return {key: visit(child) for key, child in value.items()}
        if isinstance(value, list):
            result = []
            for child in value:
                if isinstance(child, dict) and child.get('t') in ('Para', 'Plain'):
                    segment = []
                    for inline in child['c']:
                        if inline.get('t') == 'Math' and inline['c'][0]['t'] == 'DisplayMath':
                            if segment: result.append({'t': child['t'], 'c': visit(segment)}); segment = []
                            formula = inline['c'][1]
                            tags = re.findall(r'\\tag\{([^{}]+)\}', formula)
                            if len(tags) > 1: raise WordError('DUPLICATE_EQUATION_TAG')
                            labels = re.findall(r'\\label\{([^{}]+)\}', formula)
                            formula = re.sub(r'\\(?:tag|label)\{[^{}]+\}', '', formula)
                            marker = '__WMS_DISPLAY_' + str(len(equations)) + '__'
                            equations.append({'formula': formula, 'number': tags[0] if tags else None, 'marker': marker, 'labels': labels})
                            result.append({'t': 'Para', 'c': [{'t': 'Str', 'c': marker}]})
                        else: segment.append(inline)
                    if segment: result.append({'t': child['t'], 'c': visit(segment)})
                else: result.append(visit(child))
            return result
        return value
    ast = visit(ast)
    # Named author/affiliation metadata is semantic source content. Pandoc's
    # DOCX writer understands author, but silently ignores structured authors.
    def literal(value):
        if isinstance(value, dict):
            if value.get('t') in ('MetaString', 'Str'): return str(value['c'])
            if value.get('t') in ('Space', 'SoftBreak', 'LineBreak'): return ' '
            return literal(value.get('c', []))
        if isinstance(value, list): return ''.join(literal(child) for child in value)
        return str(value) if value is not None else ''
    meta = ast.get('meta', {})
    if 'authors' in meta and 'author' not in meta:
        authors = meta['authors']
        if authors.get('t') != 'MetaList': raise WordError('AUTHOR_METADATA_UNSUPPORTED')
        records = []
        for author in authors['c']:
            if author.get('t') == 'MetaMap':
                fields = author['c']; name = literal(fields.get('name'))
                details = [key + ': ' + literal(fields[key]) for key in ('affiliation', 'orcid') if fields.get(key)]
                name += ' (' + '; '.join(details) + ')' if details else ''
            else: name = literal(author)
            if not name: raise WordError('AUTHOR_METADATA_UNSUPPORTED')
            records.append({'t': 'MetaInlines', 'c': [{'t': 'Str', 'c': name}]})
        meta['author'] = {'t': 'MetaList', 'c': records}
    affiliations = meta.get('affiliations')
    if affiliations:
        if affiliations.get('t') != 'MetaList': raise WordError('AFFILIATION_METADATA_UNSUPPORTED')
        blocks = []
        for entry in affiliations['c']:
            fields = entry.get('c', {}) if entry.get('t') == 'MetaMap' else {}
            text = (literal(fields.get('index')) + ': ' if fields.get('index') else '') + literal(fields.get('name')) if fields else literal(entry)
            blocks.append({'t': 'Para', 'c': [{'t': 'Str', 'c': text}]})
        ast['blocks'] = blocks + ast['blocks']
    from .references import prepare
    reference_ledger = prepare(ast, equations)
    write_json(folder / 'references.source.json', reference_ledger)
    bibliography = ast.get('meta', {}).get('bibliography')
    if bibliography:
        def literal(value):
            if isinstance(value, dict):
                if value.get('t') in ('MetaString', 'Str'): return value['c']
                return literal(value.get('c', []))
            if isinstance(value, list): return ''.join(literal(child) for child in value)
            return ''
        name = literal(bibliography)
        file = (path.parent / name).resolve()
        if file not in permitted or not file.is_file(): raise WordError('BIBLIOGRAPHY_RESOURCE_MISSING', name)
        ast['meta']['bibliography'] = {'t': 'MetaString', 'c': str(file)}
    write_json(folder / 'source-layout-ledger.json', ignored_layout)
    write_json(folder / 'content.ast.json', ast)
    write_json(folder / 'equations.source.json', equations)
    return ast, equations


def clean(element, spec):
    styles = spec['styles']
    normal = 'Normal' if 'Normal' in styles else next(iter(styles))
    mapping = {'BodyText': normal, 'FirstParagraph': normal, 'Compact': normal,
        'ImageCaption': 'Caption', 'TableCaption': 'Caption', 'Title': 'Heading1'}
    for paragraph in element.iter(W + 'p'):
        props = paragraph.find(W + 'pPr')
        style = props.find(W + 'pStyle') if props is not None else None
        name = style.get(W + 'val') if style is not None else normal
        name = mapping.get(name, name)
        if name not in styles: name = normal
        numbering = deepcopy(props.find(W + 'numPr')) if props is not None and props.find(W + 'numPr') is not None else None
        if props is not None: paragraph.remove(props)
        props = E.Element(W + 'pPr'); paragraph.insert(0, props)
        E.SubElement(props, W + 'pStyle').set(W + 'val', name)
        if numbering is not None: props.append(numbering)
    for run in element.iter(W + 'r'):
        props = run.find(W + 'rPr')
        if props is not None:
            for child in list(props):
                if child.tag not in {W + t for t in ('b', 'i', 'vertAlign', 'strike', 'dstrike', 'rtl', 'lang')}: props.remove(child)
    for math in element.iter(M + 'r'):
        props = math.find(W + 'rPr')
        if props is not None: math.remove(props)
    if list(element.iter(W + 'sectPr')): raise WordError('CONTENT_SECTION_FORBIDDEN')
    return element


class ContentComposer(NativeComposer):
    def __init__(self, document, spec):
        super().__init__(document, preserve_styles=True); self.spec = spec
    def add_styles(self, source, element): clean(element, self.spec)
    def add_numberings(self, source, element): import_lists(source, self.doc, [element])


def compose(markdown, manifest, output, *, region=None, resources=(), data=None):
    require_actions(); check_identity(manifest)
    spec = manifest['spec']; region = region or spec['regions'][-1]['id']
    if region not in {r['id'] for r in spec['regions']}: raise WordError('CONTENT_REGION_UNKNOWN', region)
    output = Path(output); output.parent.mkdir(parents=True, exist_ok=True)
    diagnostics = output.parent / (output.stem + '-source'); diagnostics.mkdir(exist_ok=True)
    ast, equations = parse(markdown, resources, diagnostics)
    promoted_title = None
    if data and ast['blocks'] and ast['blocks'][0].get('t') == 'Header':
        def title_text(value):
            if isinstance(value, dict):
                if value.get('t') == 'Str': return value['c']
                if value.get('t') in ('Space', 'SoftBreak'): return ' '
                return title_text(value.get('c', []))
            if isinstance(value, list): return ''.join(title_text(child) for child in value)
            return ''
        title = title_text(ast['blocks'][0]['c'][2])
        if title and title in [data.get(key) for key in ('title', 'title_zh', 'title_en')]:
            promoted_title = title; ast['blocks'].pop(0)
    native = diagnostics / 'content.native.docx'
    run(['pandoc', '-f', 'json', '-t', 'docx', '-o', native, *(['--citeproc'] if ast.get('meta', {}).get('bibliography') else [])], input_text=json.dumps(ast), reject_warnings=True)
    source, destination = Document(native), Document(manifest['master'])
    expected_math = [deepcopy(m) for m in source.element.body.iter(M + 'oMath')]
    inline_attribute = '{' + NS['wms'] + '}inline_source'
    for index, math in enumerate(source.element.body.iter(M + 'oMath')): math.set(inline_attribute, str(index))
    slot = next((n for n in destination.element.body.iter(W + 'sdt')
        if n.find('w:sdtPr/w:tag', namespaces=NS) is not None
        and n.find('w:sdtPr/w:tag', namespaces=NS).get(W + 'val') == 'region.' + region), None)
    if slot is None: raise WordError('TEMPLATE_REGION_SLOT_MISSING', region)
    content = slot.find(W + 'sdtContent')
    for child in list(content): content.remove(child)
    elements = [clean(deepcopy(e), spec) for e in source.element.body if e.tag != W + 'sectPr']
    expected_text = [''.join(t.text or '' for t in p.iter(W + 't')) for e in elements for p in e.iter(W + 'p')]
    import_resources(source, destination, elements); import_lists(source, destination, elements)
    composer = ContentComposer(destination, spec)
    for element in elements:
        composer.add_footnotes(source, element)
        content.append(element)
    destination.save(output)
    package = Package(output)
    from .controls import apply as fill
    # Remaining regions accept semantic variables supplied explicitly by the caller.
    values = {r['id']: '' for r in spec['regions']}
    values.update({slot['key']: '' for r in spec['regions'] for slot in r.get('slots', [])}); values.update(data or {})
    fill(package, {'element': package.xml('word/document.xml')}, {'operation': 'fill', 'data': values})
    from .migration.math_metrics import configure
    metrics = configure({'fonts': spec['fonts']})
    from .math import apply as equation_apply
    ledger = []
    used_numbers = set()
    size = spec['styles'].get('Normal', next(iter(spec['styles'].values())))['size_pt']
    for index, equation in enumerate(equations, 1):
        marker = next((o for o in package.objects() if o['kind'] == 'paragraph' and o['text'] == equation['marker']), None)
        if marker is None: raise WordError('EQUATION_MARKER_MISSING')
        number = equation['number']
        if number is None and equation.get('labels'): number = str(index)
        if number is None and spec.get('equations', {}).get('numbering') == 'sequential': number = str(index)
        if number is not None:
            if number in used_numbers: raise WordError('DUPLICATE_EQUATION_NUMBER', number)
            used_numbers.add(number)
        ledger.append(equation_apply(package, marker, {'formula': equation['formula'], 'number': number, 'mode': 'display',
            'size_pt': size, 'number_width_twips': spec.get('equations', {}).get('number_width_twips', 600),
            'wrap': spec.get('equations', {}).get('long_formula', 'practical') != 'fail',
            'policy': spec.get('equations', {}).get('long_formula', 'practical'),
            'bookmarks': [record['bookmark'] for record in json.loads((diagnostics / 'references.source.json').read_text())['objects']
                          if record['sequence'] == 'Equation' and record['label'] in equation.get('labels', [])]}, metrics=metrics))
        marker['element'].getparent().remove(marker['element'])
    from .tables import apply as table_apply
    from .captions import apply as caption_apply
    tables = [o for o in package.objects() if o['kind'] == 'table'
        and not any(c.get(W + 'val') == 'word-module-studio-equation-layout' for c in o['element'].xpath('./w:tblPr/w:tblCaption', namespaces=NS))]
    for index, table in enumerate(tables, 1):
        previous = table['element'].getprevious()
        source_caption = previous if previous is not None and previous.tag == W + 'p' and previous.find('w:pPr/w:pStyle', namespaces=NS) is not None and previous.find('w:pPr/w:pStyle', namespaces=NS).get(W + 'val') == 'Caption' else None
        title = ''.join(source_caption.xpath('.//w:t/text()', namespaces=NS)) if source_caption is not None else None
        if title: previous.getparent().remove(previous)
        metadata_caption = table['element'].find('w:tblPr/w:tblCaption', namespaces=NS)
        if title is None and metadata_caption is not None: title = metadata_caption.get(W + 'val')
        table_apply(package, table, {'operation': 'format', 'header_rows': 1, 'three_line': True, **spec.get('table_policy', {}),
            **({'title': title, 'number': index, 'caption_position': spec['captions']['table']} if title else {})})
        if source_caption is not None:
            generated = table['element'].getprevious() if spec['captions']['table'] == 'above' else table['element'].getnext()
            from .captions import retain_inline
            retain_inline(generated, source_caption)
    for index, picture in enumerate([o for o in package.objects() if o['kind'] == 'picture'], 1):
        from .context import LayoutContext
        width = min(LayoutContext.resolve(package, picture).column_widths_twips) * 635
        for extent in picture['element'].xpath('.//*[local-name()="extent" or local-name()="ext"][@cx][@cy]'):
            old_width, old_height = int(extent.get('cx')), int(extent.get('cy'))
            if old_width > width: extent.set('cx', str(width)); extent.set('cy', str(round(old_height * width / old_width)))
        paragraph = picture['element'].getparent()
        caption = paragraph.getnext()
        if caption is not None and caption.tag == W + 'p' and caption.find('w:pPr/w:pStyle', namespaces=NS) is not None and caption.find('w:pPr/w:pStyle', namespaces=NS).get(W + 'val') == 'Caption':
            title = ''.join(caption.xpath('.//w:t/text()', namespaces=NS)); caption.getparent().remove(caption)
            caption_apply(package, {'element': paragraph}, {'operation': 'caption', 'sequence': 'Figure', 'title': title, 'number': index, 'position': spec['captions']['figure']})
            generated = paragraph.getprevious() if spec['captions']['figure'] == 'above' else paragraph.getnext()
            from .captions import retain_inline
            retain_inline(generated, caption)
    from .references import finish
    references = finish(package, json.loads((diagnostics / 'references.source.json').read_text()))
    from .math import fit_inline
    inline_elements = [math for math in package.xml('word/document.xml').iter(M + 'oMath') if math.get(inline_attribute) is not None]
    for math in inline_elements: del math.attrib[inline_attribute]
    inline = [fit_inline(package, element, size, metrics) for element in inline_elements]
    from .migration.math_layout import restore, tree
    recovered = {id(record['element']): restore(record['element'], record['changes']) for record in inline}
    from collections import Counter
    def key(math):
        value = deepcopy(math)
        for node in value.iter():
            for name in list(node.attrib):
                if name.startswith('{' + NS['wms'] + '}'): del node.attrib[name]
        return json.dumps(tree(value), ensure_ascii=False)
    required = Counter(key(math) for math in expected_math)
    actual = Counter(key(recovered.get(id(math), math)) for math in package.xml('word/document.xml').iter(M + 'oMath'))
    if required - actual: raise WordError('INLINE_OR_CAPTION_MATH_CHANGED')
    encoded_math = Counter(key(math) for math in package.xml('word/document.xml').iter(M + 'oMath'))
    inline_ledger = [{k: v for k, v in record.items() if k != 'element'} for record in inline]
    write_json(diagnostics / 'inline-math-layout-ledger.json', inline_ledger)
    package.objects(stamp=True); package.save()
    verify_locked(output, manifest)
    from .verify import relationships
    integrity = relationships(output)
    sdk = json.loads(run(['dotnet', '.tools/sdk/WordSdk.dll', 'validate', output]))
    actual = Counter(key(math) for math in Package(output).xml('word/document.xml').iter(M + 'oMath'))
    if actual != encoded_math: raise WordError('SAVED_MATH_CHANGED')
    actual_text = ''.join(Package(output).xml('word/document.xml').xpath('//w:t/text()', namespaces=NS))
    for text in expected_text:
        if text and not text.startswith('__WMS_DISPLAY_'):
            for reference in references['references']: text = text.replace(reference['marker'], reference['number'])
            if text not in actual_text: raise WordError('CONTENT_TEXT_CHANGED', text[:80])
    report = {'schema_version': 1, 'build_success': True, 'automatic_pass': True, 'visual_review_complete': False,
        'delivery_accepted': False, 'source_sha256': digest(markdown), 'output_sha256': digest(output),
        'template': {key: manifest[key] for key in ('id', 'version', 'spec_sha256', 'master_sha256')},
        'display_equations': len(equations), 'equations': ledger, 'references': references,
        'inline_math_preserved': len(expected_math), 'inline_math_layout': inline_ledger,
        'source_title_promoted': promoted_title, 'relationships': integrity, 'sdk': sdk}
    write_json(output.with_suffix('.compose.json'), report)
    return report
